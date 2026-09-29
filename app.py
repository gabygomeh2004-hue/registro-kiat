import streamlit as st
import json
import os
from datetime import datetime, timedelta, time
from pathlib import Path
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import uuid

# ============================================================
# CONFIGURACIÓN
# ============================================================
DATA_FILE = Path(__file__).parent / "registros.json"
INGENIEROS = [
    "Andres Buvoli",
    "Daniela Benitez",
    "Gabriela Gómez",
    "Ivan Alviz",
    "Ivan Gómez",
    "Jesús Solorzano",
    "Oscar Diaz",
]

st.set_page_config(
    page_title="Registro KIAT - Gobernación de Sucre",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ============================================================
# CONEXIÓN A GOOGLE SHEETS (para versión online compartida)
# ============================================================
def get_gsheet():
    """Intenta conectar a Google Sheets usando secrets de Streamlit.
    Retorna el worksheet o None si no está configurado."""
    try:
        if "gcp_service_account" not in st.secrets:
            return None
        import gspread
        from google.oauth2.service_account import Credentials

        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        credentials = Credentials.from_service_account_info(
            st.secrets["gcp_service_account"], scopes=scopes
        )
        gc = gspread.authorize(credentials)
        # El ID de la hoja se pone en secrets también
        sheet_id = st.secrets.get("sheet_id", "")
        if not sheet_id:
            return None
        sh = gc.open_by_key(sheet_id)
        # Usamos la primera hoja
        return sh.sheet1
    except Exception as e:
        st.sidebar.warning(f"No se pudo conectar a Google Sheets: {e}")
        return None


def load_data():
    """Carga los registros. Prioridad: Google Sheets > archivo local."""
    ws = get_gsheet()
    if ws is not None:
        try:
            records = ws.get_all_records()
            # Convertir a lista de dicts con los campos esperados
            data = []
            for r in records:
                data.append({
                    "id": str(r.get("id", "")),
                    "nombre": r.get("nombre", ""),
                    "tipo": r.get("tipo", ""),
                    "fecha": r.get("fecha", ""),
                    "hora": r.get("hora", ""),
                    "destino": r.get("destino", ""),
                    "observacion": r.get("observacion", ""),
                    "timestamp": r.get("timestamp", ""),
                })
            return data
        except Exception:
            return []

    # Fallback: archivo local
    if DATA_FILE.exists():
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []


def save_data(data):
    """Guarda los registros. Prioridad: Google Sheets > archivo local."""
    ws = get_gsheet()
    if ws is not None:
        try:
            # Limpiar y reescribir toda la hoja (simple y seguro para pocos registros)
            headers = ["id", "nombre", "tipo", "fecha", "hora", "destino", "observacion", "timestamp"]
            rows = [headers]
            for r in data:
                rows.append([
                    r.get("id", ""),
                    r.get("nombre", ""),
                    r.get("tipo", ""),
                    r.get("fecha", ""),
                    r.get("hora", ""),
                    r.get("destino", ""),
                    r.get("observacion", ""),
                    r.get("timestamp", ""),
                ])
            ws.clear()
            ws.update("A1", rows)
            return
        except Exception as e:
            st.error(f"Error al guardar en Google Sheets: {e}")
            return

    # Fallback: archivo local
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_now():
    """Hora actual en Colombia (UTC-5)"""
    return datetime.now(ZoneInfo("America/Bogota"))


def determinar_jornada(ahora=None):
    if ahora is None:
        ahora = get_now()
    hora = ahora.time()
    dia = ahora.weekday()  # 0=lunes ... 5=sábado, 6=domingo

    if dia == 5:  # sábado
        return "sabado"
    if hora >= time(8, 0) and hora < time(14, 0):
        return "manana"
    if hora >= time(14, 0):
        return "tarde"
    return "fuera"  # antes de las 8am


def tipo_llegada_label(tipo):
    return {
        "llegada_manana": "Llegada Mañana",
        "llegada_tarde": "Llegada Tarde",
        "llegada_sabado": "Llegada Sábado",
        "salida_campo": "Salida de Campo",
    }.get(tipo, tipo)


# ============================================================
# GENERAR DOCX
# ============================================================
def set_cell_shading(cell, color_hex):
    """Aplica color de fondo a una celda"""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def generar_reporte_semanal(registros, fecha_inicio, fecha_fin):
    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    # Encabezado
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("LABORATORIO KIAT")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 70, 127)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitulo.add_run("Gobernación de Sucre")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    titulo2 = doc.add_paragraph()
    titulo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo2.add_run("REGISTRO SEMANAL DE LLEGADAS Y SALIDAS DE CAMPO")
    run.bold = True
    run.font.size = Pt(13)

    periodo = doc.add_paragraph()
    periodo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = periodo.add_run(
        f"Período: {fecha_inicio.strftime('%d/%m/%Y')} – {fecha_fin.strftime('%d/%m/%Y')}"
    )
    run.font.size = Pt(11)

    doc.add_paragraph()

    # Filtrar registros de la semana
    registros_semana = []
    for r in registros:
        try:
            f = datetime.strptime(r["fecha"], "%Y-%m-%d").date()
            if fecha_inicio <= f <= fecha_fin:
                registros_semana.append(r)
        except Exception:
            continue

    registros_semana.sort(key=lambda x: (x["fecha"], x["hora"]))

    # --- Tabla de Llegadas ---
    doc.add_heading("1. Llegadas al Laboratorio", level=1)

    llegadas = [r for r in registros_semana if r["tipo"].startswith("llegada")]
    if not llegadas:
        doc.add_paragraph("No se registraron llegadas en este período.")
    else:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Fecha", "Hora", "Nombre", "Jornada", "Observación"]
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)
            set_cell_shading(hdr_cells[i], "00467F")
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)

        for r in llegadas:
            row = table.add_row().cells
            row[0].text = datetime.strptime(r["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")
            row[1].text = r["hora"]
            row[2].text = r["nombre"]
            row[3].text = tipo_llegada_label(r["tipo"])
            row[4].text = r.get("observacion", "") or "—"
            for cell in row:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

    doc.add_paragraph()

    # --- Tabla de Salidas de Campo ---
    doc.add_heading("2. Salidas de Campo", level=1)

    salidas = [r for r in registros_semana if r["tipo"] == "salida_campo"]
    if not salidas:
        doc.add_paragraph("No se registraron salidas de campo en este período.")
    else:
        table2 = doc.add_table(rows=1, cols=5)
        table2.style = "Table Grid"
        table2.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers2 = ["Fecha", "Hora", "Nombre", "Destino / Motivo", "Observación"]
        hdr_cells2 = table2.rows[0].cells
        for i, h in enumerate(headers2):
            hdr_cells2[i].text = h
            for paragraph in hdr_cells2[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)
            set_cell_shading(hdr_cells2[i], "2E7D32")
            for paragraph in hdr_cells2[i].paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)

        for r in salidas:
            row = table2.add_row().cells
            row[0].text = datetime.strptime(r["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")
            row[1].text = r["hora"]
            row[2].text = r["nombre"]
            row[3].text = r.get("destino", "") or "—"
            row[4].text = r.get("observacion", "") or "—"
            for cell in row:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

    doc.add_paragraph()
    doc.add_paragraph()

    # Pie de página / firmas
    firmas = doc.add_paragraph()
    firmas.add_run("_______________________________\n").font.size = Pt(10)
    firmas.add_run("Responsable del Laboratorio\n").font.size = Pt(9)
    firmas.add_run("Laboratorio KIAT – Gobernación de Sucre").font.size = Pt(8)

    # Resumen
    doc.add_paragraph()
    resumen = doc.add_paragraph()
    resumen.add_run("Resumen del período:\n").bold = True
    resumen.add_run(f"• Total llegadas: {len(llegadas)}\n")
    resumen.add_run(f"• Total salidas de campo: {len(salidas)}\n")
    resumen.add_run(f"• Generado el: {get_now().strftime('%d/%m/%Y %H:%M')}")

    # Guardar
    nombre_archivo = f"Registro_KIAT_{fecha_inicio.strftime('%Y%m%d')}_{fecha_fin.strftime('%Y%m%d')}.docx"
    ruta = Path(__file__).parent / nombre_archivo
    doc.save(ruta)
    return ruta, len(llegadas), len(salidas)


# ============================================================
# INTERFAZ PRINCIPAL
# ============================================================
def main():
    st.markdown(
        """
        <style>
        .main-header {
            background: linear-gradient(90deg, #00467F 0%, #0077B6 100%);
            padding: 1rem 1.5rem;
            border-radius: 10px;
            color: white;
            margin-bottom: 1.5rem;
        }
        .stButton>button {
            width: 100%;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        """
        <div class="main-header">
            <h1 style="margin:0; color:white;">🔬 Laboratorio KIAT</h1>
            <p style="margin:0; opacity:0.9;">Gobernación de Sucre — Registro de Llegadas y Salidas de Campo</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    data = load_data()

    # Sidebar
    with st.sidebar:
        st.header("Menú")
        seccion = st.radio(
            "Seleccione una opción:",
            [
                "Registrar Llegada",
                "Registrar Salida de Campo",
                "Ver Registros del Día",
                "Historial / Filtros",
                "Generar Reporte Semanal",
                "Administrar Registros",
            ],
            label_visibility="collapsed",
        )
        st.divider()
        ahora = get_now()
        st.caption(f"🕐 Hora actual: **{ahora.strftime('%H:%M:%S')}**")
        st.caption(f"📅 Fecha: **{ahora.strftime('%d/%m/%Y')}**")
        jornada_actual = determinar_jornada(ahora)
        if jornada_actual == "manana":
            st.success("Jornada: Mañana (desde 8:00 a.m.)")
        elif jornada_actual == "tarde":
            st.info("Jornada: Tarde (desde 2:00 p.m.)")
        elif jornada_actual == "sabado":
            st.warning("Hoy es Sábado")
        else:
            st.error("Fuera de jornada laboral")

        # Indicador de modo de almacenamiento
        if get_gsheet() is not None:
            st.success("☁️ Modo compartido (Google Sheets)")
        else:
            st.info("💾 Modo local (este computador)")

    # ============================================================
    # 1. REGISTRAR LLEGADA
    # ============================================================
    if seccion == "Registrar Llegada":
        st.subheader("📥 Registrar Llegada")

        ahora = get_now()
        jornada = determinar_jornada(ahora)

        col1, col2 = st.columns([2, 1])

        with col1:
            nombre = st.selectbox("Seleccione su nombre:", INGENIEROS, key="llegada_nombre")

        with col2:
            if jornada == "sabado":
                tipo_default = "llegada_sabado"
            elif jornada == "manana":
                tipo_default = "llegada_manana"
            elif jornada == "tarde":
                tipo_default = "llegada_tarde"
            else:
                tipo_default = "llegada_manana"

            tipo_opciones = {
                "llegada_manana": "Llegada Mañana (desde 8:00 a.m.)",
                "llegada_tarde": "Llegada Tarde (desde 2:00 p.m.)",
                "llegada_sabado": "Llegada Sábado (esporádica)",
            }
            tipo = st.selectbox(
                "Tipo de llegada:",
                list(tipo_opciones.keys()),
                format_func=lambda x: tipo_opciones[x],
                index=list(tipo_opciones.keys()).index(tipo_default) if tipo_default in tipo_opciones else 0,
            )

        observacion = st.text_input("Observación (opcional):", placeholder="Ej: Llegó un poco antes por cita...")

        if jornada == "fuera" and tipo != "llegada_sabado":
            st.warning("⚠️ Está registrando fuera del horario habitual de la jornada seleccionada.")

        if st.button("✅ Registrar Llegada", type="primary", use_container_width=True):
            hoy = ahora.strftime("%Y-%m-%d")
            existe = any(
                r["nombre"] == nombre
                and r["fecha"] == hoy
                and r["tipo"] == tipo
                for r in data
            )
            if existe:
                st.error(f"Ya existe un registro de {tipo_llegada_label(tipo)} para {nombre} hoy.")
            else:
                nuevo = {
                    "id": str(uuid.uuid4())[:8],
                    "nombre": nombre,
                    "tipo": tipo,
                    "fecha": hoy,
                    "hora": ahora.strftime("%H:%M:%S"),
                    "destino": "",
                    "observacion": observacion.strip() if observacion else "",
                    "timestamp": ahora.isoformat(),
                }
                data.append(nuevo)
                save_data(data)
                st.success(
                    f"✅ Llegada registrada: **{nombre}** — {tipo_llegada_label(tipo)} "
                    f"a las **{nuevo['hora']}** del {ahora.strftime('%d/%m/%Y')}"
                )
                st.balloons()
                st.rerun()

    # ============================================================
    # 2. REGISTRAR SALIDA DE CAMPO
    # ============================================================
    elif seccion == "Registrar Salida de Campo":
        st.subheader("🚗 Registrar Salida de Campo")

        ahora = get_now()

        nombre = st.selectbox("Seleccione su nombre:", INGENIEROS, key="salida_nombre")
        destino = st.text_input(
            "Destino o motivo (recomendado):",
            placeholder="Ej: Inspección en Sincelejo / Muestreo en Tolú...",
        )
        observacion = st.text_input("Observación adicional (opcional):")

        if st.button("✅ Registrar Salida de Campo", type="primary", use_container_width=True):
            nuevo = {
                "id": str(uuid.uuid4())[:8],
                "nombre": nombre,
                "tipo": "salida_campo",
                "fecha": ahora.strftime("%Y-%m-%d"),
                "hora": ahora.strftime("%H:%M:%S"),
                "destino": destino.strip() if destino else "",
                "observacion": observacion.strip() if observacion else "",
                "timestamp": ahora.isoformat(),
            }
            data.append(nuevo)
            save_data(data)
            st.success(
                f"✅ Salida de campo registrada: **{nombre}** a las **{nuevo['hora']}** "
                f"— Destino: {nuevo['destino'] or 'No especificado'}"
            )
            st.rerun()

    # ============================================================
    # 3. VER REGISTROS DEL DÍA
    # ============================================================
    elif seccion == "Ver Registros del Día":
        st.subheader("📋 Registros del Día Actual")

        hoy = get_now().strftime("%Y-%m-%d")
        registros_hoy = [r for r in data if r["fecha"] == hoy]
        registros_hoy.sort(key=lambda x: x["hora"])

        if not registros_hoy:
            st.info("No hay registros para el día de hoy.")
        else:
            st.write(f"**Total registros hoy:** {len(registros_hoy)}")

            llegadas_hoy = [r for r in registros_hoy if r["tipo"].startswith("llegada")]
            salidas_hoy = [r for r in registros_hoy if r["tipo"] == "salida_campo"]

            if llegadas_hoy:
                st.markdown("#### Llegadas")
                for r in llegadas_hoy:
                    st.markdown(
                        f"- **{r['hora']}** — {r['nombre']} ({tipo_llegada_label(r['tipo'])})"
                        + (f" · _{r['observacion']}_" if r.get("observacion") else "")
                    )

            if salidas_hoy:
                st.markdown("#### Salidas de Campo")
                for r in salidas_hoy:
                    dest = r.get("destino") or "Sin destino"
                    st.markdown(
                        f"- **{r['hora']}** — {r['nombre']} → {dest}"
                        + (f" · _{r['observacion']}_" if r.get("observacion") else "")
                    )

    # ============================================================
    # 4. HISTORIAL / FILTROS
    # ============================================================
    elif seccion == "Historial / Filtros":
        st.subheader("🔍 Historial y Filtros")

        col1, col2, col3 = st.columns(3)
        with col1:
            filtro_nombre = st.selectbox(
                "Persona:",
                ["Todas"] + INGENIEROS,
                key="filtro_nombre",
            )
        with col2:
            filtro_tipo = st.selectbox(
                "Tipo:",
                [
                    "Todos",
                    "llegada_manana",
                    "llegada_tarde",
                    "llegada_sabado",
                    "salida_campo",
                ],
                format_func=lambda x: "Todos" if x == "Todos" else tipo_llegada_label(x),
            )
        with col3:
            rango = st.date_input(
                "Rango de fechas:",
                value=(get_now().date() - timedelta(days=7), get_now().date()),
                key="filtro_fechas",
            )

        filtrados = data.copy()
        if filtro_nombre != "Todas":
            filtrados = [r for r in filtrados if r["nombre"] == filtro_nombre]
        if filtro_tipo != "Todos":
            filtrados = [r for r in filtrados if r["tipo"] == filtro_tipo]

        if isinstance(rango, (list, tuple)) and len(rango) == 2:
            f_ini, f_fin = rango
            filtrados = [
                r
                for r in filtrados
                if f_ini <= datetime.strptime(r["fecha"], "%Y-%m-%d").date() <= f_fin
            ]

        filtrados.sort(key=lambda x: (x["fecha"], x["hora"]), reverse=True)

        st.write(f"**Resultados:** {len(filtrados)} registros")

        if filtrados:
            tabla = []
            for r in filtrados:
                tabla.append(
                    {
                        "Fecha": datetime.strptime(r["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y"),
                        "Hora": r["hora"],
                        "Nombre": r["nombre"],
                        "Tipo": tipo_llegada_label(r["tipo"]),
                        "Destino/Obs.": r.get("destino") or r.get("observacion") or "—",
                    }
                )
            st.dataframe(tabla, use_container_width=True, hide_index=True)

    # ============================================================
    # 5. GENERAR REPORTE SEMANAL
    # ============================================================
    elif seccion == "Generar Reporte Semanal":
        st.subheader("📄 Generar Reporte Semanal (.docx)")

        st.markdown(
            "Seleccione el rango de fechas de la semana que desea incluir en el reporte."
        )

        hoy = get_now().date()
        inicio_semana = hoy - timedelta(days=hoy.weekday())
        fin_semana = inicio_semana + timedelta(days=6)

        col1, col2 = st.columns(2)
        with col1:
            fecha_inicio = st.date_input("Desde:", value=inicio_semana)
        with col2:
            fecha_fin = st.date_input("Hasta:", value=fin_semana)

        if fecha_inicio > fecha_fin:
            st.error("La fecha de inicio no puede ser posterior a la fecha de fin.")
        else:
            if st.button("📥 Generar documento Word", type="primary", use_container_width=True):
                with st.spinner("Generando documento..."):
                    ruta, n_llegadas, n_salidas = generar_reporte_semanal(
                        data, fecha_inicio, fecha_fin
                    )
                st.success(
                    f"✅ Documento generado correctamente.\n\n"
                    f"- Llegadas incluidas: **{n_llegadas}**\n"
                    f"- Salidas de campo: **{n_salidas}**"
                )
                with open(ruta, "rb") as f:
                    st.download_button(
                        label="⬇️ Descargar archivo .docx",
                        data=f,
                        file_name=ruta.name,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True,
                    )

    # ============================================================
    # 6. ADMINISTRAR REGISTROS (editar / anular)
    # ============================================================
    elif seccion == "Administrar Registros":
        st.subheader("⚙️ Administrar Registros")
        st.caption("Permite anular o eliminar registros erróneos. Use con precaución.")

        if not data:
            st.info("No hay registros aún.")
        else:
            data_ordenada = sorted(data, key=lambda x: x.get("timestamp", ""), reverse=True)[:40]

            opciones = {
                f"{r['id']} | {r['fecha']} {r['hora']} — {r['nombre']} ({tipo_llegada_label(r['tipo'])})": r["id"]
                for r in data_ordenada
            }

            seleccionado = st.selectbox(
                "Seleccione el registro a eliminar:",
                list(opciones.keys()),
            )

            if st.button("🗑️ Eliminar registro seleccionado", type="secondary"):
                id_eliminar = opciones[seleccionado]
                data = [r for r in data if r["id"] != id_eliminar]
                save_data(data)
                st.success("Registro eliminado correctamente.")
                st.rerun()

    # Footer
    st.divider()
    st.caption(
        "Laboratorio KIAT — Gobernación de Sucre | Sistema de registro de llegadas y salidas de campo"
    )


if __name__ == "__main__":
    main()
